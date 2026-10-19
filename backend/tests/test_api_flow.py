"""Fluxo ponta a ponta: cadastro → pesquisa → fila → verificação → filtros → exportação → histórico."""

import asyncio
import io

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import app
from app.providers import registry
from app.providers.base import WebResult
from app.services import search_runner
from app.services.dedup import ingest_record
from app.db import SessionLocal
from app.models import Business, Job, Organization
from app.worker import run_worker
from tests.fakes import FakeChecker, FakePlaces, FakeWeb, rec

H = {"X-Requested-With": "prospecia"}


@pytest.fixture
def client():
    with TestClient(app) as c:
        yield c


def register(c, email="ana@ex.com"):
    r = c.post("/api/auth/register", json={"organization": "Agência Web", "name": "Ana", "email": email,
                                           "password": "senha-forte-1"})
    assert r.status_code == 201, r.text
    return r.json()


def drain():
    async def go():
        await run_worker(asyncio.Event(), concurrency=2, idle_exit=True)
    asyncio.run(go())


@pytest.fixture
def fake_sources(monkeypatch):
    places = FakePlaces([
        [rec("p1", "Barbearia do Zé", phone="+55 41 3333-4444", lat=-25.43, lng=-49.27, rating=4.7, reviews_count=80,
             business_status="operational"),
         rec("p2", "Barber Prime", phone="+55 41 3333-5555", website="https://barberprime.com.br")],
        [rec("p3", "Corte Fino Barbearia", phone="+55 41 3333-6666"),
         rec("p1", "Barbearia do Zé", phone="+55 41 3333-4444")],  # duplicada na página seguinte
    ])
    registry.register_place(places)
    web = FakeWeb({"Corte Fino": [WebResult("https://cortefino.com.br", "Corte Fino Barbearia", "", 1)]})
    monkeypatch.setattr(search_runner, "web_provider", lambda preferred=None: web)
    checker = FakeChecker({
        "barberprime.com.br": {"title": "Barber Prime", "text": "Agende: (41) 3333-5555 Curitiba"},
        "cortefino.com.br": {"title": "Corte Fino Barbearia", "text": "(41) 3333-6666"},
    })
    import app.worker as worker_mod
    monkeypatch.setattr(worker_mod, "SiteChecker", lambda: checker)
    yield places, web


SEARCH = {"locations": [{"city": "Curitiba", "state": "PR", "country": "Brasil"}], "categories": ["barbearias"],
          "sources": ["fake_places"], "max_results": 50, "web_verification": True}


def test_fluxo_completo(client, fake_sources):
    register(client)
    # CSRF: escrita via cookie exige cabeçalho
    assert client.post("/api/searches", json=SEARCH).status_code == 403

    est = client.post("/api/searches/estimate", json=SEARCH, headers=H).json()
    assert est["total_usd"] > 0 and "budget" in est

    s = client.post("/api/searches", json=SEARCH, headers=H).json()
    assert s["status"] == "queued"
    drain()

    detail = client.get(f"/api/searches/{s['id']}").json()
    assert detail["status"] == "completed", detail
    assert detail["total_found"] == 3  # duplicata não contada
    assert detail["total_verified"] == detail["total_to_verify"] == 3
    assert detail["status_counts"] == {"SEM_SITE_IDENTIFICADO": 1, "POSSUI_SITE": 2}

    page = client.get("/api/businesses", params={"search_id": s["id"], "digital_status": "SEM_SITE_IDENTIFICADO"}).json()
    assert page["total"] == 1
    ze = page["items"][0]
    assert ze["name"] == "Barbearia do Zé" and ze["source_website"] is None and ze["confidence"] >= 60

    d = client.get(f"/api/businesses/{ze['id']}").json()
    assert d["verifications"][0]["evidence"] and d["verifications"][0]["sources_consulted"]
    assert d["source_records"][0]["source"] == "fake_places"

    # CRM: notas e status
    client.post(f"/api/businesses/{ze['id']}/interactions", json={"kind": "note", "content": "Ligar segunda"}, headers=H)
    client.patch(f"/api/businesses/{ze['id']}", json={"lead_status": "contato_realizado"}, headers=H)
    d = client.get(f"/api/businesses/{ze['id']}").json()
    assert d["lead_status"] == "contato_realizado" and len(d["interactions"]) == 2

    # Filtros combinados e ordenação
    r = client.get("/api/businesses", params={"has_phone": True, "min_confidence": 50, "sort": "name", "order": "asc"}).json()
    assert [i["name"] for i in r["items"]] == sorted(i["name"] for i in r["items"])

    # Exportações
    csv = client.get("/api/exports/businesses.csv", params={"search_id": s["id"]})
    assert csv.status_code == 200 and "Barbearia do Zé" in csv.content.decode("utf-8-sig")
    assert "Status da presença digital" in csv.content.decode("utf-8-sig")
    xl = client.get("/api/exports/businesses.xlsx", params={"digital_status": "POSSUI_SITE"})
    wb = load_workbook(io.BytesIO(xl.content))
    assert wb["Leads"].max_row == 3 and "Sobre" in wb.sheetnames

    # Dashboard
    dash = client.get("/api/dashboard").json()
    assert dash["total"] == 3 and dash["by_status"]["POSSUI_SITE"] == 2

    # Reabrir pesquisa: histórico preservado e nova execução não duplica empresas
    again = client.post(f"/api/searches/{s['id']}/rerun", headers=H).json()
    drain()
    d2 = client.get(f"/api/searches/{again['id']}").json()
    assert d2["status"] == "completed" and d2["total_new"] == 0 and d2["total_found"] == 3
    assert d2["total_to_verify"] == 0  # verificações recentes reaproveitadas
    assert client.get("/api/businesses").json()["total"] == 3
    assert len(client.get("/api/searches").json()) == 2


def test_permissoes_e_limite(client, fake_sources):
    register(client)
    client.post("/api/users", json={"name": "Beto", "email": "beto@ex.com", "password": "senha-forte-2",
                                    "role": "member", "monthly_budget_usd": 0}, headers=H)
    with TestClient(app) as other:
        other.post("/api/auth/login", json={"email": "beto@ex.com", "password": "senha-forte-2"})
        r = other.post("/api/searches", json=SEARCH, headers=H)
        assert r.status_code == 402  # limite de consumo zerado
        with SessionLocal() as db:
            b, _ = ingest_record(db, 1, rec("x1", "Empresa X"))
            db.commit()
        assert other.delete(f"/api/businesses/{b.id}", headers=H).status_code == 403
    assert client.delete(f"/api/businesses/{b.id}", headers=H).status_code == 204


def test_isolamento_entre_organizacoes(client, fake_sources):
    register(client)
    with SessionLocal() as db:
        b, _ = ingest_record(db, 1, rec("x1", "Empresa X"))
        db.commit()
    with TestClient(app) as other:
        register(other, "outra@ex.com")
        assert other.get(f"/api/businesses/{b.id}").status_code == 404
        assert other.get("/api/businesses").json()["total"] == 0


def test_falha_temporaria_e_retomada(client, monkeypatch, fake_sources):
    places, _ = fake_sources
    places.fail_times = 1
    monkeypatch.setattr("app.services.queue.timedelta", lambda seconds=0, minutes=0: __import__("datetime").timedelta(0))
    register(client)
    s = client.post("/api/searches", json=SEARCH, headers=H).json()
    drain()
    d = client.get(f"/api/searches/{s['id']}").json()
    assert d["status"] == "completed" and d["total_found"] == 3
    assert any(e["level"] == "error" for e in d["events"])  # erro de integração registrado


def test_falha_persistente_marca_pesquisa_como_falha(client, monkeypatch, fake_sources):
    places, _ = fake_sources
    places.fail_times = 99
    monkeypatch.setattr("app.services.queue.timedelta", lambda seconds=0, minutes=0: __import__("datetime").timedelta(0))
    register(client)
    s = client.post("/api/searches", json=SEARCH, headers=H).json()
    drain()
    d = client.get(f"/api/searches/{s['id']}").json()
    assert d["status"] == "failed" and "Retomar" in d["error_message"]
    places.fail_times = 0
    client.post(f"/api/searches/{s['id']}/resume", headers=H)
    drain()
    assert client.get(f"/api/searches/{s['id']}").json()["status"] == "completed"


def test_pausa_e_retomada(client, fake_sources):
    register(client)
    s = client.post("/api/searches", json=SEARCH, headers=H).json()
    client.post(f"/api/searches/{s['id']}/pause", headers=H)
    drain()
    assert client.get(f"/api/searches/{s['id']}").json()["status"] == "paused"
    client.post(f"/api/searches/{s['id']}/resume", headers=H)
    drain()
    assert client.get(f"/api/searches/{s['id']}").json()["status"] == "completed"


def test_deduplicacao_entre_fontes():
    with SessionLocal() as db:
        db.add(Organization(id=1, name="Org"))
        db.flush()
        a, new_a = ingest_record(db, 1, rec("g1", "Padaria São José", phone="(41) 3333-1111", lat=-25.4, lng=-49.2,
                                            source="google_places"))
        b, new_b = ingest_record(db, 1, rec("node/9", "Padaria Sao Jose LTDA", phone="+55 41 3333-1111",
                                            website="https://padariasaojose.com.br", source="osm"))
        c, new_c = ingest_record(db, 1, rec("g2", "Padaria São José", phone="(11) 9999-0000", city="São Paulo",
                                            lat=-23.5, lng=-46.6, source="google_places"))
        db.commit()
        assert new_a and not new_b and new_c
        assert a.id == b.id != c.id
        assert sorted(a.sources) == ["google_places", "osm"]
        assert a.source_website == "https://padariasaojose.com.br"
        assert db.query(Business).count() == 2


def test_documentacao_openapi(client):
    assert "/api/searches" in client.get("/openapi.json").json()["paths"]


def test_fila_nao_duplica_verificacao(client, fake_sources):
    register(client)
    with SessionLocal() as db:
        b, _ = ingest_record(db, 1, rec("x1", "Empresa X"))
        db.commit()
    client.post(f"/api/businesses/{b.id}/verify", headers=H)
    r = client.post(f"/api/businesses/{b.id}/verify", headers=H).json()
    assert r["queued"] is False
    with SessionLocal() as db:
        assert db.query(Job).filter(Job.kind == "verify_business").count() == 1
