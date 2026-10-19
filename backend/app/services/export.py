"""Exportação CSV/Excel preservando links para as fontes."""

from __future__ import annotations

import csv
import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from app.models import Business
from app.verification.classifier import STATUS_LABELS

LEAD_LABELS = {
    "novo": "Novo lead", "revisado": "Lead revisado", "contato_realizado": "Contato realizado",
    "aguardando_resposta": "Aguardando resposta", "interessado": "Interessado", "proposta_enviada": "Proposta enviada",
    "convertido": "Cliente convertido", "sem_interesse": "Sem interesse", "nao_contatar": "Não contatar",
}

COLUMNS = [
    ("ID", lambda b: b.id),
    ("Empresa", lambda b: b.name),
    ("Categoria", lambda b: b.category),
    ("Endereço", lambda b: b.address),
    ("Bairro", lambda b: b.neighborhood),
    ("Cidade", lambda b: b.city),
    ("Estado", lambda b: b.state),
    ("CEP", lambda b: b.postal_code),
    ("Telefone comercial", lambda b: b.phone),
    ("E-mail comercial público", lambda b: b.email),
    ("Perfil Maps / fonte", lambda b: b.maps_url),
    ("Site informado na fonte", lambda b: b.source_website),
    ("Site oficial identificado", lambda b: b.official_website),
    ("Domínio candidato (não confirmado)", lambda b: b.candidate_website),
    ("Redes sociais", lambda b: " | ".join(b.social_links or [])),
    ("Status da presença digital", lambda b: STATUS_LABELS.get(b.digital_status, b.digital_status)),
    ("Confiança (%)", lambda b: b.confidence),
    ("Justificativa", lambda b: b.classification_reason),
    ("Apenas redes sociais", lambda b: "Sim" if b.social_only else "Não"),
    ("Avaliação", lambda b: b.rating),
    ("Nº de avaliações", lambda b: b.reviews_count),
    ("Pontuação do lead", lambda b: b.lead_score),
    ("Status da prospecção", lambda b: LEAD_LABELS.get(b.lead_status, b.lead_status)),
    ("Fontes", lambda b: ", ".join(b.sources or [])),
    ("Última verificação", lambda b: b.last_verified_at.strftime("%Y-%m-%d %H:%M") if b.last_verified_at else None),
    ("Dado fictício (demonstração)", lambda b: "SIM" if b.is_fictitious else "Não"),
]
URL_COLUMNS = {"Perfil Maps / fonte", "Site informado na fonte", "Site oficial identificado", "Domínio candidato (não confirmado)"}
NOTE = ("Atenção: 'Site não informado no Maps' significa apenas que o cadastro de origem não apresenta site; "
        "não confirma ausência de site na internet. A pontuação é uma prioridade estimada, não garantia de interesse. "
        "Use os contatos conforme a LGPD e respeite o status 'Não contatar'.")


def _safe(v):
    # Evita injeção de fórmulas em planilhas (CSV/Excel).
    if isinstance(v, str) and v[:1] in ("=", "+", "-", "@"):
        return "'" + v
    return v


def to_csv(rows: list[Business]) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";")
    w.writerow([c[0] for c in COLUMNS])
    for b in rows:
        w.writerow([_safe(f(b)) if f(b) is not None else "" for _, f in COLUMNS])
    return ("﻿" + buf.getvalue()).encode("utf-8")


def to_xlsx(rows: list[Business]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Leads"
    ws.append([c[0] for c in COLUMNS])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1E3A8A")
    for b in rows:
        ws.append([_safe(f(b)) for _, f in COLUMNS])
        r = ws.max_row
        for idx, (name, _) in enumerate(COLUMNS, start=1):
            if name in URL_COLUMNS:
                cell = ws.cell(row=r, column=idx)
                if isinstance(cell.value, str) and cell.value.startswith("http"):
                    cell.hyperlink = cell.value
                    cell.font = Font(color="1D4ED8", underline="single")
    for idx, (name, _) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = 45 if name in ("Justificativa", "Endereço") else 22
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    info = wb.create_sheet("Sobre")
    info.append(["Gerado em", datetime.now().strftime("%Y-%m-%d %H:%M")])
    info.append(["Total de registros", len(rows)])
    info.append(["Observações", NOTE])
    info.column_dimensions["B"].width = 120
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
